"""Classeur Excel à trier : exporter les résultats, les trier dans Excel, réimporter.

Trier les pages avec des cases à cocher dans la page de recherche n'était pas tenable au-delà de
quelques dizaines de résultats. Le tri se fait maintenant dans Excel, et le classeur réimporté
devient une sélection où chercher ensuite. Les invariants tenus ici : les pages gardées sont
exactement les lignes à « oui », quoi qu'on ait fait du classeur entre-temps (lignes triées ou
supprimées, colonnes déplacées) ; une valeur non reconnue n'est jamais prise pour un « oui » et
se voit dans le bilan ; un classeur ne se réimporte pas dans un autre index ; la chaîne des
recherches voyage avec le fichier. Et la langue : le classeur suit celle de l'interface, se
relit quelle que soit la sienne, et un refus est un code que la page traduit.
"""
import io
import json
from urllib.parse import unquote

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import Workbook, load_workbook

import results_workbook
from results_workbook import WorkbookError, read_kept_pages
from selections_service import SelectionsService
from services import IndexesService

from conftest import build, make_collection

REGISTRES = {"REGA": {"REGA_1": ["chat", "noir"], "REGA_2": ["chat", "blanc"]},
             "REGB": {"REGB_1": ["chat", "gris"], "REGB_2": ["merle"]}}
FILTRES = dict(year_from=None, year_to=None, fuzzy_threshold=None, whole_word=False, exclude=None)
ID_COL = 9      # colonne « Identifiant »


@pytest.fixture
def corpus(data_dir):
    IndexesService._vocab_cache.clear()
    make_collection(data_dir, "COL", REGISTRES)
    for reg, names in {"REGA": ["REGA_1.jpg", "REGA_2.jpg"], "REGB": ["REGB_1.jpg"]}.items():
        folder = data_dir / "collections" / "COL" / "scans" / reg
        folder.mkdir(parents=True, exist_ok=True)
        for name in names:
            (folder / name).write_bytes(b"image")
    build("idx1")
    return data_dir


def routes():
    from routers import indexes
    return indexes


def exporter(q="chat", **kw):
    return routes().export_results_xlsx("idx1", q=q, **{**FILTRES, "within": None, "lang": "fr", **kw})


def importer(data, filename="tri.xlsx", index_id="idx1"):
    return routes().import_selection(index_id, UploadFile(io.BytesIO(data), filename=filename))


def affichee(page_name):
    return IndexesService._display_page_name(page_name)


def trier(data, decisions=None, supprimer=(), inverser=False):
    """Ce qu'un utilisateur ferait dans Excel : changer des « Garder », supprimer des lignes,
    trier. `decisions` et `supprimer` désignent les pages par leur nom affiché."""
    wb = load_workbook(io.BytesIO(data))
    ws = wb["Pages"]
    lignes = [[c.value for c in row] for row in ws.iter_rows(min_row=2)]
    lignes = [l for l in lignes if affichee(l[ID_COL - 1]) not in supprimer]
    for l in lignes:
        if affichee(l[ID_COL - 1]) in (decisions or {}):
            l[0] = decisions[affichee(l[ID_COL - 1])]
    if inverser:
        lignes.reverse()
    ws.delete_rows(2, ws.max_row)
    for l in lignes:
        ws.append(l)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def classeur_minimal(lignes, entetes=("Garder", "Identifiant (ne pas modifier)"), feuille="Pages"):
    wb = Workbook()
    ws = wb.active
    ws.title = feuille
    ws.append(list(entetes))
    for l in lignes:
        ws.append(list(l))
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


# ── L'export ─────────────────────────────────────────────────────────────────

def test_le_classeur_a_trier(corpus):
    reponse = exporter()

    assert reponse.media_type == results_workbook.MEDIA_TYPE
    assert unquote(reponse.headers["content-disposition"].split("UTF-8''")[1]) == "chat.xlsx"
    wb = load_workbook(io.BytesIO(reponse.body))
    assert wb.sheetnames == ["Pages", "Recherche", "archimed"]
    assert wb["archimed"].sheet_state == "hidden"

    ws = wb["Pages"]
    assert [c.value for c in ws[1]] == results_workbook.LABELS["fr"]["columns"]
    lignes = list(ws.iter_rows(min_row=2, values_only=True))
    assert {affichee(l[ID_COL - 1]): l[0] for l in lignes} == {
        "REGA_1": "oui", "REGA_2": "oui", "REGB_1": "oui"}
    assert {l[2]: l[4] for l in lignes}["REGA_1"] == "chat × 1"
    deux_termes = load_workbook(io.BytesIO(exporter(q="chat noir").body))["Pages"]
    assert deux_termes.cell(row=2, column=5).value == "chat × 1, noir × 1"
    assert ws.freeze_panes == "A2" and ws.auto_filter.ref == "A1:I4"
    validation = ws.data_validations.dataValidation[0]
    assert validation.formula1 == '"oui,non"' and str(validation.sqref) == "A2:A4"
    # Un lien vers l'image de chaque page, pour décider en la regardant.
    lien = ws.cell(row=2, column=8).hyperlink.target
    assert lien.startswith("file:///") and lien.endswith(".jpg")

    meta = {r[0]: r[1] for r in wb["archimed"].iter_rows(values_only=True)}
    assert meta["index_id"] == "idx1"
    assert [(s["query"], s["found"]) for s in json.loads(meta["steps"])] == [("chat", 3)]


def test_le_classeur_parle_la_langue_de_l_interface(corpus):
    """Une interface en anglais exporte un classeur en anglais : en-têtes, oui/non, feuille
    de recherche et nom du fichier."""
    reponse = exporter(q="chat", exclude="gris", lang="en")

    assert unquote(reponse.headers["content-disposition"].split("UTF-8''")[1]) == "chat without gris.xlsx"
    wb = load_workbook(io.BytesIO(reponse.body))
    assert wb.sheetnames == ["Pages", "Search", "archimed"]
    ws = wb["Pages"]
    assert [c.value for c in ws[1]] == results_workbook.LABELS["en"]["columns"]
    assert {l[0] for l in ws.iter_rows(min_row=2, values_only=True)} == {"yes"}
    assert ws.data_validations.dataValidation[0].formula1 == '"yes,no"'
    assert ("Searched in", "the whole index") in list(wb["Search"].iter_rows(values_only=True, max_col=2))


# ── L'aller-retour ───────────────────────────────────────────────────────────

def test_on_garde_exactement_les_oui_quoi_qu_on_ait_fait_du_classeur(corpus):
    trie = trier(exporter().body, decisions={"REGA_2": "non"}, supprimer={"REGB_1"}, inverser=True)

    reponse = importer(trie, filename="chats noirs.xlsx")

    sel = reponse["selection"]
    assert (reponse["kept"], reponse["dropped"], reponse["ignored"]) == (1, 1, 0)
    assert sel["name"] == "chats noirs" and sel["source_file"] == "chats noirs.xlsx"
    assert [(s["query"], s["found"], s["kept"]) for s in sel["steps"]] == [("chat", 3, 1)]
    # Et la recherche suivante se fait dedans.
    dedans = routes().search_in_index("idx1", q="chat", within=sel["id"], **FILTRES)
    assert {affichee(p["page_name"]) for p in dedans["pages"]} == {"REGA_1"}


def test_la_chaine_voyage_avec_le_classeur(corpus):
    premier = importer(trier(exporter().body, decisions={"REGB_1": "non"}), filename="chats.xlsx")

    second = exporter(q="noir", within=premier["selection"]["id"])

    assert unquote(second.headers["content-disposition"].split("UTF-8''")[1]) == "chat › noir.xlsx"
    wb = load_workbook(io.BytesIO(second.body))
    assert ("Cherché dans", "la sélection « chats »") in list(wb["Recherche"].iter_rows(values_only=True, max_col=2))
    sel = importer(second.body, filename="chat › noir.xlsx")["selection"]
    assert [(s["query"], s["found"], s["kept"]) for s in sel["steps"]] == [
        ("chat", 3, 2), ("noir", 1, 1)]


def test_un_classeur_en_anglais_se_reimporte(corpus):
    """Exporté sur un poste en anglais (« Keep », yes / no), relu sans y penser."""
    trie = trier(exporter(lang="en").body, decisions={"REGA_2": "no"}, supprimer={"REGB_1"})

    reponse = importer(trie)

    assert (reponse["kept"], reponse["dropped"], reponse["ignored"]) == (1, 1, 0)
    assert {affichee(p) for p in SelectionsService.pages("idx1", reponse["selection"]["id"])} == {"REGA_1"}


# ── La lecture de « Garder » ─────────────────────────────────────────────────

def test_les_valeurs_de_garder(corpus):
    lu = read_kept_pages(classeur_minimal([
        ("Oui ", "p1"), ("x", "p2"), (1, "p3"), (True, "p4"),
        ("NON", "p5"), (None, "p6"),
        ("peut-être", "p7"),
        ("oui", None),              # ligne sans identifiant : ignorée sans bruit
        ("oui", "p1"),              # doublon
    ]))

    assert lu["pages"] == ["p1", "p2", "p3", "p4"]
    assert (lu["kept"], lu["dropped"], lu["ignored"]) == (4, 2, 1)
    assert lu["index_id"] is None and lu["steps"] == []


def test_les_colonnes_sont_reperees_par_leur_entete(corpus):
    lu = read_kept_pages(classeur_minimal([("p1", "note", "oui"), ("p2", "", "non")],
                                          entetes=("Identifiant", "Note", "garder")))

    assert lu["pages"] == ["p1"]


def code_de(data):
    with pytest.raises(WorkbookError) as refus:
        read_kept_pages(data)
    return refus.value.code, refus.value.params


def test_un_classeur_refuse_dit_pourquoi_par_un_code(corpus):
    """Des codes, que la page traduit dans sa langue, plutôt que des phrases en français."""
    # Un autre classeur, sans rien de commun avec une liste ArchiMED (l'inventaire des registres…).
    assert code_de(classeur_minimal([("HPC-1", "Registre 1")], entetes=("Cote", "Titre"),
                                    feuille="Inventaire")) == ("not_archimed", {"sheet": "Inventaire"})
    # Une liste ArchiMED dont on a supprimé une des deux colonnes lues.
    assert code_de(classeur_minimal([("p1",)], entetes=("Identifiant",)))[0] == "missing_keep_column"
    assert code_de(classeur_minimal([("yes",)], entetes=("Keep",)))[0] == "missing_id_column"
    assert code_de(b"Garder;Identifiant\noui;p1\n")[0] == "unreadable"


def test_les_refus_de_l_import(corpus):
    build("idx2", name="Autre")
    venu_de_idx1 = exporter().body

    with pytest.raises(HTTPException) as autre_index:
        importer(venu_de_idx1, index_id="idx2")
    assert autre_index.value.status_code == 400
    assert (autre_index.value.detail["code"], autre_index.value.detail["params"]) == (
        "other_index", {"index": "Idx"})

    with pytest.raises(HTTPException) as rien:
        importer(trier(venu_de_idx1, decisions={"REGA_1": "non", "REGA_2": "non", "REGB_1": "non"}))
    assert (rien.value.status_code, rien.value.detail["code"]) == (400, "nothing_kept")

    with pytest.raises(HTTPException) as illisible:
        importer(b"pas un classeur")
    assert (illisible.value.status_code, illisible.value.detail["code"]) == (400, "unreadable")

    with pytest.raises(HTTPException) as inconnu:
        importer(venu_de_idx1, index_id="inconnu")
    assert inconnu.value.status_code == 404
