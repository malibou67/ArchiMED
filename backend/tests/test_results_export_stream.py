"""Export CSV / Excel en flux : dire où en est l'export, puis livrer le fichier.

Exporter la liste à trier était un simple lien : le navigateur attendait en silence que le
serveur ait relancé la recherche puis inventorié les images registre par registre (près d'une
minute sur le partage pour une recherche large), et une erreur ne s'affichait nulle part. Le
flux raconte ces étapes, et le fichier arrive dans son dernier événement. Les invariants tenus
ici : les étapes arrivent dans l'ordre, l'inventaire compte les registres ; le fichier livré est
exactement celui du téléchargement direct ; une sélection inconnue reste un vrai 404.
"""
import asyncio
import base64
import io
import json

import pytest
from fastapi import HTTPException
from openpyxl import load_workbook

from services import IndexesService

from conftest import build, make_collection

REGISTRES = {"REGA": {"REGA_1": ["chat", "noir"], "REGA_2": ["chat", "blanc"]},
             "REGB": {"REGB_1": ["chat", "gris"], "REGB_2": ["merle"]}}
FILTRES = dict(year_from=None, year_to=None, fuzzy_threshold=None, whole_word=False, exclude=None,
               within=None, lang="fr")


@pytest.fixture
def corpus(data_dir):
    IndexesService._vocab_cache.clear()
    make_collection(data_dir, "COL", REGISTRES)
    for reg, names in {"REGA": ["REGA_1.jpg", "REGA_2.jpg"], "REGB": ["REGB_1.jpg"]}.items():
        folder = data_dir / "collections" / "COL" / "scans" / reg
        folder.mkdir(parents=True, exist_ok=True)
        for name in names:
            (folder / name).write_bytes(b"image")
    build("idx1")
    return data_dir


def routes():
    from routers import indexes
    return indexes


def flux(fmt, index_id="idx1", q="chat", **kw):
    """Vide la réponse en flux comme le ferait Starlette, et rend ses événements."""
    reponse = routes().export_results_stream(index_id, fmt=fmt, q=q, **{**FILTRES, **kw})

    async def vider():
        return b"".join([m if isinstance(m, bytes) else m.encode() async for m in reponse.body_iterator])

    return [json.loads(l) for l in asyncio.run(vider()).decode().splitlines() if l.strip()]


def phases(events):
    vues = [e['phase'] for e in events if e['type'] == 'progress']
    return [p for i, p in enumerate(vues) if i == 0 or p != vues[i - 1]]


def test_l_export_excel_raconte_ses_etapes_puis_livre_le_classeur(corpus):
    events = flux("xlsx")

    # La recherche, l'inventaire (un événement par registre), l'écriture, puis le fichier.
    assert phases(events)[-4:] == ['scan', 'build', 'resolve', 'write']
    inventaire = [(e['current'], e['total']) for e in events if e.get('phase') == 'resolve']
    assert inventaire == [(0, 2), (1, 2), (2, 2)]
    fichier = events[-1]
    assert fichier['type'] == 'file' and fichier['filename'] == "chat.xlsx"

    # Le même classeur que le téléchargement direct.
    livre = load_workbook(io.BytesIO(base64.b64decode(fichier['data'])))["Pages"]
    direct = load_workbook(io.BytesIO(routes().export_results_xlsx("idx1", q="chat", **FILTRES).body))["Pages"]
    assert list(livre.iter_rows(values_only=True)) == list(direct.iter_rows(values_only=True))


def test_l_export_csv_livre_le_meme_fichier_que_le_telechargement_direct(corpus):
    events = flux("csv")

    fichier = events[-1]
    assert (fichier['filename'], fichier['media_type']) == ("idx1_resultats.csv", "text/csv")
    sans_lang = {k: v for k, v in FILTRES.items() if k != 'lang'}
    assert base64.b64decode(fichier['data']) == routes().export_results_csv("idx1", q="chat", **sans_lang).body


def test_les_erreurs_de_l_export_en_flux(corpus):
    # Une sélection inconnue : un vrai 404, avant d'ouvrir le flux.
    with pytest.raises(HTTPException) as inconnue:
        flux("xlsx", within="sel_20990101000000_000000")
    assert inconnue.value.status_code == 404
    # Un index inconnu : le flux le dit.
    assert flux("xlsx", index_id="inconnu")[-1]['type'] == 'error'
