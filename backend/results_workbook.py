"""Classeur Excel des résultats d'une recherche : à trier, puis à réimporter.

Trier des centaines de pages avec des cases à cocher dans la page de recherche n'était pas
tenable. Le tri se fait dans Excel : une ligne par page, une colonne « Garder » pré-remplie à
« oui » (liste déroulante oui / non), les filtres d'Excel pour écarter un registre ou un mot d'un
coup, et un lien vers l'image de chaque page pour décider en la regardant. Réimporté, le classeur
devient une sélection dans laquelle chercher (voir `selections_service`).

Le classeur est rédigé dans la langue de l'interface (« Garder » / oui, ou « Keep » / yes), et
l'import lit indifféremment les deux : un fichier exporté sur un poste en anglais se réimporte
sur un poste en français. Les erreurs de lecture sont des codes (`WorkbookError`) que la page
traduit, plutôt que des phrases dans une seule langue.

À l'import, seules deux colonnes comptent, repérées par leur en-tête : « Garder » et
« Identifiant », le nom interne de la page. Trier, filtrer, déplacer des colonnes ou supprimer
des lignes (une ligne supprimée est une page écartée) reste donc sans effet sur la lecture. Une
feuille masquée garde l'index d'origine et la chaîne des recherches : un fichier ne se réimporte
pas dans un autre index, et la sélection sait d'où elle vient.

Ce module ne fait que du format : ni recherche, ni accès aux données.
"""
import io
import json
import re
from datetime import datetime
from typing import Any, Dict, List, Optional
from zipfile import BadZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.worksheet.datavalidation import DataValidation

FORMAT_VERSION = 1
PAGES_SHEET = "Pages"          # même nom dans les deux langues
META_SHEET = "archimed"
MEDIA_TYPE = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
LANGS = ("fr", "en")

# Tout le texte du classeur, par langue. Les colonnes sont dans l'ordre de la feuille « Pages ».
LABELS: Dict[str, Dict[str, Any]] = {
    "fr": {
        "columns": ["Garder", "Registre", "Page", "Occurrences", "Mots trouvés", "Collection",
                    "Modèle", "Image", "Identifiant (ne pas modifier)"],
        "yes": "oui", "no": "non",
        "prompt": "oui : garder la page ; non : l'écarter",
        "error": "Choisir oui ou non.",
        "search_sheet": "Recherche",
        "exported": "Export ArchiMED", "date": "%d/%m/%Y à %H:%M", "index": "Index",
        "searched_in": "Cherché dans", "in_selection": "la sélection « {name} »",
        "whole_index": "tout l'index", "found": "Pages trouvées",
        "steps": ["Étape", "Recherche", "Mots exclus", "Mot entier", "Ressemblance", "Période",
                  "Pages trouvées", "Pages gardées"],
        "exact": "exacte", "to_sort": "à trier", "without": "sans",
        "howto_title": "Mode d'emploi",
        "howto": "Dans la feuille « Pages », mettre « non » dans la colonne Garder pour écarter une "
                 "page (ou supprimer sa ligne). Réimporter ensuite ce fichier depuis la page "
                 "Recherche d'ArchiMED (« Importer un Excel ») pour chercher dans les pages gardées.",
        "default_name": "resultats",
    },
    "en": {
        "columns": ["Keep", "Register", "Page", "Occurrences", "Words found", "Collection",
                    "Model", "Image", "Identifier (do not edit)"],
        "yes": "yes", "no": "no",
        "prompt": "yes: keep the page; no: drop it",
        "error": "Choose yes or no.",
        "search_sheet": "Search",
        "exported": "ArchiMED export", "date": "%Y-%m-%d %H:%M", "index": "Index",
        "searched_in": "Searched in", "in_selection": "the selection \"{name}\"",
        "whole_index": "the whole index", "found": "Pages found",
        "steps": ["Step", "Search", "Excluded words", "Whole word", "Similarity", "Period",
                  "Pages found", "Pages kept"],
        "exact": "exact", "to_sort": "to sort", "without": "without",
        "howto_title": "How to",
        "howto": "In the \"Pages\" sheet, put \"no\" in the Keep column to drop a page (or delete "
                 "its row). Then import this file back from the ArchiMED Search page (\"Import an "
                 "Excel file\") to search within the pages kept.",
        "default_name": "results",
    },
}
_WIDTHS = [9, 24, 28, 12, 40, 18, 16, 55, 34]

# Ce que l'import reconnaît, quelle que soit la langue du classeur.
KEEP_HEADERS = {LABELS[lang]["columns"][0].lower() for lang in LANGS}       # garder, keep
ID_HEADER_PREFIXES = ("identifiant", "identifier")
# Valeurs de « Garder ». La liste déroulante n'offre que oui / non (yes / no), mais une valeur
# collée, ou saisie dans un autre tableur, échappe à la validation.
YES = {"oui", "o", "x", "1", "vrai", "yes", "y", "true"}
NO = {"non", "n", "0", "faux", "no", "false", ""}

_HEADER_FONT = Font(bold=True)
_HEADER_FILL = PatternFill(start_color="E8EEF6", end_color="E8EEF6", fill_type="solid")
_DROPPED_FONT = Font(color="9E9E9E")
_DROPPED_FILL = PatternFill(start_color="F2F2F2", end_color="F2F2F2", fill_type="solid")
_FORBIDDEN_IN_FILENAME = re.compile(r'[<>:"/\\|?*\x00-\x1f]')


class WorkbookError(ValueError):
    """Classeur refusé à l'import. `code` (et `params`) sont traduits par la page ; le message
    français reste pour les journaux et les appels directs à l'API."""

    def __init__(self, code: str, message: str, **params):
        super().__init__(message)
        self.code = code
        self.params = params


def _lang(lang: Optional[str]) -> str:
    return lang if lang in LANGS else "fr"


def step_label(step: Dict[str, Any], lang: str = "fr") -> str:
    """« cancer sans sein, poumon » : une étape de l'entonnoir, en une ligne."""
    label = step.get('query') or ''
    if step.get('excluded_terms'):
        label += f" {LABELS[_lang(lang)]['without']} {', '.join(step['excluded_terms'])}"
    return label


def export_filename(chain: List[Dict[str, Any]], lang: str = "fr") -> str:
    """Nom du fichier proposé au téléchargement : la chaîne des recherches. C'est aussi le nom
    que prendra la sélection une fois le fichier réimporté, à moins de le renommer."""
    stem = _FORBIDDEN_IN_FILENAME.sub('', ' › '.join(step_label(s, lang) for s in chain)).strip(' .')
    return f"{(stem or LABELS[_lang(lang)]['default_name'])[:120]}.xlsx"


def _words_label(words: Dict[str, List[str]]) -> str:
    return ', '.join(f"{w} × {len(occs)}" for w, occs in sorted(words.items()))


def build_results_workbook(index_id: str, index_name: str, rows: List[Dict[str, Any]],
                           chain: List[Dict[str, Any]], within_name: Optional[str] = None,
                           lang: str = "fr") -> bytes:
    """Le classeur d'une recherche, dans la langue `lang` (fr / en). `rows` vient de
    `IndexesService.resolve_result_pages` ; `chain` est la suite des recherches depuis tout
    l'index, la dernière étant celle-ci."""
    L = LABELS[_lang(lang)]
    wb = Workbook()

    # ── Pages ────────────────────────────────────────────────────────────
    ws = wb.active
    ws.title = PAGES_SHEET
    ws.append(L["columns"])
    for col, width in enumerate(_WIDTHS, start=1):
        cell = ws.cell(row=1, column=col)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        ws.column_dimensions[get_column_letter(col)].width = width
    image_col = L["columns"].index("Image") + 1
    for r in rows:
        ws.append([L["yes"], r.get('registre') or '', r.get('page') or '', r.get('occurrences') or 0,
                   _words_label(r.get('words') or {}), r.get('source') or '', r.get('model') or '',
                   r.get('chemin') or '', r.get('page_name') or ''])
        if r.get('path'):
            # Lien vers l'image : on décide en regardant la page, sans quitter Excel.
            cell = ws.cell(row=ws.max_row, column=image_col)
            cell.hyperlink = r['path'].resolve().as_uri()
            cell.style = "Hyperlink"

    last_row = len(rows) + 1
    last_col = get_column_letter(len(L["columns"]))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{last_col}{last_row}"
    if rows:
        keep = DataValidation(type="list", formula1=f'"{L["yes"]},{L["no"]}"', allow_blank=True)
        keep.promptTitle = keep.errorTitle = L["columns"][0]
        keep.prompt = L["prompt"]
        keep.error = L["error"]
        keep.showInputMessage = True
        keep.showErrorMessage = True
        ws.add_data_validation(keep)
        keep.add(f"A2:A{last_row}")
        # Une ligne écartée se grise : on voit d'un coup d'œil ce qui reste.
        ws.conditional_formatting.add(
            f"A2:{last_col}{last_row}",
            FormulaRule(formula=[f'$A2="{L["no"]}"'], font=_DROPPED_FONT, fill=_DROPPED_FILL))

    # ── Recherche (pour un lecteur humain) ───────────────────────────────
    info = wb.create_sheet(L["search_sheet"])
    info.column_dimensions['A'].width = 16
    info.column_dimensions['B'].width = 40
    for col in "CDEFGH":
        info.column_dimensions[col].width = 18
    index_label = index_name if index_name == index_id else f"{index_name} ({index_id})"
    for label, value in [
        (L["exported"], datetime.now().strftime(L["date"])),
        (L["index"], index_label),
        (L["searched_in"], L["in_selection"].format(name=within_name) if within_name else L["whole_index"]),
        (L["found"], len(rows)),
    ]:
        info.append([label, value])
        info.cell(row=info.max_row, column=1).font = _HEADER_FONT
    info.append([])
    info.append(L["steps"])
    for cell in info[info.max_row]:
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
    for i, step in enumerate(chain, start=1):
        period = ""
        if step.get('year_from') is not None or step.get('year_to') is not None:
            period = f"{step.get('year_from') or '…'} – {step.get('year_to') or '…'}"
        fuzzy = step.get('fuzzy_threshold')
        info.append([
            i, step.get('query') or '', ', '.join(step.get('excluded_terms') or []),
            L["yes"] if step.get('whole_word') else L["no"],
            f"≥ {fuzzy} %" if fuzzy is not None and fuzzy < 100 else L["exact"],
            period, step.get('found'),
            step.get('kept') if i < len(chain) else L["to_sort"],
        ])
    info.append([])
    info.append([L["howto_title"], L["howto"]])
    info.cell(row=info.max_row, column=1).font = _HEADER_FONT
    info.cell(row=info.max_row, column=2).alignment = Alignment(wrap_text=True, vertical="top")

    # ── Feuille masquée : de quoi réimporter ─────────────────────────────
    meta = wb.create_sheet(META_SHEET)
    meta.append(["format", FORMAT_VERSION])
    meta.append(["index_id", index_id])
    meta.append(["index_name", index_name])
    meta.append(["steps", json.dumps(chain, ensure_ascii=False)])
    meta.sheet_state = "hidden"

    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def _clean_steps(raw: Any) -> List[Dict[str, Any]]:
    """Étapes relues dans la feuille masquée, réduites aux champs connus. Une feuille abîmée ne
    fait pas échouer l'import : la sélection n'aura simplement pas de chaîne."""
    if not isinstance(raw, list):
        return []
    steps = []
    for s in raw:
        if not isinstance(s, dict) or not isinstance(s.get('query'), str):
            return []
        try:
            steps.append({
                'query': s['query'],
                'whole_word': bool(s.get('whole_word')),
                'fuzzy_threshold': int(s['fuzzy_threshold']) if s.get('fuzzy_threshold') is not None else None,
                'excluded_terms': [str(t) for t in s.get('excluded_terms') or []],
                'year_from': int(s['year_from']) if s.get('year_from') is not None else None,
                'year_to': int(s['year_to']) if s.get('year_to') is not None else None,
                'found': int(s.get('found') or 0),
                'kept': int(s['kept']) if s.get('kept') is not None else None,
            })
        except (TypeError, ValueError):
            return []
    return steps


def _keep_value(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, bool):
        return "oui" if value else "non"
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip().lower()


def read_kept_pages(data: bytes) -> Dict[str, Any]:
    """Relit un classeur trié, en français comme en anglais. Renvoie `{pages, kept, dropped,
    ignored, index_id, index_name, steps}` : les identifiants des pages à « oui », dans l'ordre
    du fichier, et le bilan. Lève `WorkbookError` si le fichier ne peut pas être lu :
    `unreadable` (pas un .xlsx), `not_archimed` (un classeur, mais pas une liste exportée par
    ArchiMED), `missing_keep_column` / `missing_id_column` (une liste ArchiMED dont on a
    supprimé une colonne nécessaire)."""
    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, KeyError, OSError, ValueError) as e:
        raise WorkbookError("unreadable", "Ce fichier n'est pas un classeur Excel (.xlsx) lisible.") from e
    try:
        is_archimed = META_SHEET in wb.sheetnames
        ws = wb[PAGES_SHEET] if PAGES_SHEET in wb.sheetnames else wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(v).strip().lower() if v is not None else '' for v in next(rows, ())]
        keep_col = next((i for i, h in enumerate(header) if h in KEEP_HEADERS), None)
        id_col = next((i for i, h in enumerate(header) if h.startswith(ID_HEADER_PREFIXES)), None)
        if keep_col is None and id_col is None and not is_archimed:
            # Ni nos colonnes ni notre feuille masquée : un autre classeur (un inventaire…).
            raise WorkbookError(
                "not_archimed",
                f"Ce classeur (feuille « {ws.title} ») n'est pas une liste exportée par ArchiMED.",
                sheet=ws.title)
        if keep_col is None:
            raise WorkbookError("missing_keep_column",
                                f"Colonne « Garder » introuvable dans la feuille « {ws.title} ».",
                                sheet=ws.title)
        if id_col is None:
            raise WorkbookError("missing_id_column",
                                f"Colonne « Identifiant » introuvable dans la feuille « {ws.title} ».",
                                sheet=ws.title)

        pages: List[str] = []
        dropped = ignored = 0
        for row in rows:
            page = row[id_col] if id_col < len(row) else None
            if page is None or not str(page).strip():
                continue            # ligne vide
            keep = _keep_value(row[keep_col] if keep_col < len(row) else None)
            if keep in YES:
                pages.append(str(page).strip())
            elif keep in NO:
                dropped += 1
            else:
                ignored += 1

        meta: Dict[str, Any] = {}
        if is_archimed:
            for row in wb[META_SHEET].iter_rows(values_only=True):
                if row and row[0] is not None and len(row) > 1:
                    meta[str(row[0])] = row[1]
    finally:
        wb.close()

    try:
        steps = _clean_steps(json.loads(meta['steps'])) if meta.get('steps') else []
    except ValueError:
        steps = []
    pages = list(dict.fromkeys(pages))
    return {
        'pages': pages,
        'kept': len(pages),
        'dropped': dropped,
        'ignored': ignored,
        'index_id': str(meta['index_id']) if meta.get('index_id') else None,
        'index_name': str(meta['index_name']) if meta.get('index_name') else None,
        'steps': steps,
    }
